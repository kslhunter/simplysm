#!/usr/bin/env python3
"""Extract text and images from PPTX files with per-slide coordinates."""

import sys
import io
import subprocess
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")


def ensure_packages():
    packages = {"python-pptx": "pptx"}
    for pip_name, import_name in packages.items():
        try:
            __import__(import_name)
        except ImportError:
            print(f"Installing package: {pip_name}...", file=sys.stderr)
            subprocess.check_call([sys.executable, "-m", "pip", "install", pip_name],
                                  stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def emu_to_inches(emu):
    """Convert EMU to inches (1 decimal place)."""
    if emu is None:
        return "?"
    return f"{emu / 914400:.1f}"


def extract(file_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(file_path)
    stem = Path(file_path).stem
    out_dir = Path(file_path).parent / f"{stem}_files"
    img_idx = 0

    print(f"# {Path(file_path).name}\n")

    for slide_num, slide in enumerate(prs.slides, 1):
        print(f"## Slide {slide_num}\n")

        for shape in slide.shapes:
            left = emu_to_inches(shape.left)
            top = emu_to_inches(shape.top)
            pos = f"(left={left}\", top={top}\")"

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_idx += 1
                blob = shape.image.blob
                content_type = shape.image.content_type
                ext = content_type.split("/")[-1].replace("jpeg", "jpg")
                out_dir.mkdir(parents=True, exist_ok=True)
                img_path = out_dir / f"img_{img_idx:03d}.{ext}"
                img_path.write_bytes(blob)
                print(f"[IMG] {pos} {img_path}")

            elif hasattr(shape, "text") and shape.text.strip():
                text = shape.text.strip().replace("\n", "\n       ")
                print(f"[TXT] {pos} {text}")

        print()

    if img_idx > 0:
        print(f"---\n{img_idx} image(s) saved: {out_dir}")
    else:
        print("---\nNo images")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_pptx.py <file.pptx>", file=sys.stderr)
        sys.exit(1)
    ensure_packages()
    extract(sys.argv[1])
